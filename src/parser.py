import os
import re
import hashlib
import json
import zipfile
import shutil
import xml.etree.ElementTree as ET
from pathlib import Path
from bs4 import BeautifulSoup
import pymupdf4llm

def compute_file_hash(filepath):
    sha = hashlib.sha256()
    with open(filepath, 'rb') as f:
        while chunk := f.read(65536):
            sha.update(chunk)
    return sha.hexdigest()

def parse_pptx_to_markdown(filepath):
    chunks = []
    full_md_list = []

    try:
        import pptx
        prs = pptx.Presentation(str(filepath))

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_lines = []
            slide_title = None

            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title_text = slide.shapes.title.text.strip()
                if title_text:
                    slide_title = title_text
                    slide_lines.append(f'# Slide {slide_idx}: {title_text}')
                    slide_lines.append('')

            if not slide_title:
                slide_lines.append(f'# Slide {slide_idx}')
                slide_lines.append('')

            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue

                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        p_text = paragraph.text.strip()
                        if not p_text:
                            continue
                        indent = '  ' * paragraph.level
                        if paragraph.level > 0:
                            slide_lines.append(f'{indent}- {p_text}')
                        elif p_text.startswith(('#', '-', '*', '>')):
                            slide_lines.append(p_text)
                        else:
                            slide_lines.append(f'- {p_text}')

                elif shape.has_table:
                    table = shape.table
                    table_lines = []
                    if len(table.rows) > 0:
                        headers = [cell.text.strip().replace(chr(10), ' ') or '-' for cell in table.rows[0].cells]
                        table_lines.append('| ' + ' | '.join(headers) + ' |')
                        table_lines.append('| ' + ' | '.join(['---'] * len(headers)) + ' |')

                        for row in table.rows[1:]:
                            row_cells = [cell.text.strip().replace(chr(10), ' ') or '-' for cell in row.cells]
                            table_lines.append('| ' + ' | '.join(row_cells) + ' |')

                        slide_lines.append('')
                        slide_lines.extend(table_lines)
                        slide_lines.append('')

            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_lines.append('')
                    slide_lines.append(f'> **Speaker Notes:** {notes}')

            slide_text = chr(10).join(slide_lines).strip()
            if slide_text:
                chunks.append({
                    'page': slide_idx,
                    'text': slide_text
                })
                full_md_list.append(f'<!-- Page {slide_idx} -->' + chr(10) + slide_text)

    except Exception as pptx_err:
        print(f'[Parser Notice] python-pptx parser fallback triggered for {Path(filepath).name}: {pptx_err}')
        try:
            with zipfile.ZipFile(filepath, 'r') as zf:
                slide_files = sorted(
                    [f for f in zf.namelist() if f.startswith('ppt/slides/slide') and f.endswith('.xml')],
                    key=lambda x: int(re.search(r'slide(\d+)\.xml', x).group(1)) if re.search(r'slide(\d+)\.xml', x) else 0
                )
                for slide_idx, sf in enumerate(slide_files, start=1):
                    xml_data = zf.read(sf)
                    tree = ET.fromstring(xml_data)
                    texts = [elem.text for elem in tree.iter() if elem.text and elem.tag.endswith('t')]
                    if texts:
                        clean_text = chr(10).join([f'- {t.strip()}' for t in texts if t.strip()])
                        slide_body = f'# Slide {slide_idx}' + chr(10) + chr(10) + clean_text
                        chunks.append({'page': slide_idx, 'text': slide_body})
                        full_md_list.append(f'<!-- Page {slide_idx} -->' + chr(10) + slide_body)
        except Exception as e:
            print(f'[Parser Error] Failed XML extraction of PPTX {Path(filepath).name}: {e}')

    combined_md = (chr(10) + chr(10) + '---' + chr(10) + chr(10)).join(full_md_list)
    return chunks, combined_md


def parse_docx_to_markdown(filepath):
    chunks = []
    lines = []
    try:
        with zipfile.ZipFile(filepath, 'r') as zf:
            if 'word/document.xml' in zf.namelist():
                xml_data = zf.read('word/document.xml')
                tree = ET.fromstring(xml_data)
                for elem in tree.iter():
                    if elem.tag.endswith('p'):
                        p_texts = [node.text for node in elem.iter() if node.text and node.tag.endswith('t')]
                        p_str = ''.join(p_texts).strip()
                        if p_str:
                            lines.append(p_str)
        text = (chr(10) + chr(10)).join(lines).strip()
        if text:
            chunks.append({'page': 1, 'text': text})
        return chunks, text
    except Exception as e:
        print(f'[Parser Error] Failed to parse DOCX {Path(filepath).name}: {e}')
        return [], ''


class DocumentParser:
    def __init__(self, output_dir, parsed_dir):
        self.output_dir = Path(output_dir)
        self.parsed_dir = Path(parsed_dir)
        self.parsed_dir.mkdir(parents=True, exist_ok=True)
        self.hash_cache_file = self.parsed_dir / '.file_hashes.json'
        self.hashes = self._load_hashes()

    def _load_hashes(self):
        if self.hash_cache_file.exists():
            try:
                with open(self.hash_cache_file, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except Exception:
                return {}
        return {}

    def _save_hashes(self):
        with open(self.hash_cache_file, 'w', encoding='utf-8') as f:
            json.dump(self.hashes, f, indent=2)

    def parse_file(self, filepath, force=False):
        filepath = Path(filepath).resolve()
        if not filepath.exists() or filepath.name.startswith('.'):
            return []

        try:
            rel_path = filepath.relative_to(self.output_dir)
        except ValueError:
            rel_path = Path(filepath.name)

        parts = rel_path.parts
        sem_label = parts[0] if len(parts) > 0 else 'Unknown_Semester'
        course_name = parts[1] if len(parts) > 1 else 'Unknown_Course'

        curr_hash = compute_file_hash(filepath)
        cached_hash = self.hashes.get(str(rel_path))

        # Check if already up-to-date
        md_dest = self.parsed_dir / rel_path.with_suffix('.md')
        if not force and cached_hash == curr_hash and md_dest.exists():
            try:
                with open(md_dest, 'r', encoding='utf-8') as f:
                    full_md = f.read()
                
                page_sections = re.split(r'<!-- Page \d+ -->', full_md)
                if len(page_sections) > 1:
                    cached_chunks = []
                    page_matches = re.findall(r'<!-- Page (\d+) -->', full_md)
                    for i, p_num_str in enumerate(page_matches):
                        p_num = int(p_num_str)
                        if i + 1 < len(page_sections):
                            p_txt = page_sections[i+1].strip().rstrip('-').strip()
                            if p_txt:
                                cached_chunks.append({
                                    'text': p_txt,
                                    'metadata': {
                                        'source': str(filepath),
                                        'filename': filepath.name,
                                        'course': course_name,
                                        'semester': sem_label,
                                        'page': p_num
                                    }
                                })
                    if cached_chunks:
                        return cached_chunks

                return [{
                    'text': full_md,
                    'metadata': {
                        'source': str(filepath),
                        'filename': filepath.name,
                        'course': course_name,
                        'semester': sem_label,
                        'page': 1
                    }
                }]
            except Exception:
                pass

        ext = filepath.suffix.lower()
        chunks = []

        # 1. PDF Documents
        if ext == '.pdf':
            try:
                pages_data = pymupdf4llm.to_markdown(str(filepath), page_chunks=True)
                full_md_list = []
                for p in pages_data:
                    p_text = p.get('text', '').strip()
                    p_meta = p.get('metadata', {})
                    page_num = p_meta.get('page_number', 1)
                    if p_text:
                        chunks.append({
                            'text': p_text,
                            'metadata': {
                                'source': str(filepath),
                                'filename': filepath.name,
                                'course': course_name,
                                'semester': sem_label,
                                'page': page_num
                            }
                        })
                        full_md_list.append(f'<!-- Page {page_num} -->' + chr(10) + p_text)
                
                md_dest.parent.mkdir(parents=True, exist_ok=True)
                with open(md_dest, 'w', encoding='utf-8') as f:
                    f.write((chr(10) + chr(10) + '---' + chr(10) + chr(10)).join(full_md_list))
                
            except Exception as e:
                print(f'[Parser Error] Failed to parse PDF {filepath.name}: {e}')
                return []

        # 2. PowerPoint Presentations (.pptx, .ppt)
        elif ext in ('.pptx', '.ppt'):
            try:
                raw_chunks, combined_md = parse_pptx_to_markdown(filepath)
                for ch in raw_chunks:
                    chunks.append({
                        'text': ch['text'],
                        'metadata': {
                            'source': str(filepath),
                            'filename': filepath.name,
                            'course': course_name,
                            'semester': sem_label,
                            'page': ch['page']
                        }
                    })
                if combined_md:
                    md_dest.parent.mkdir(parents=True, exist_ok=True)
                    with open(md_dest, 'w', encoding='utf-8') as f:
                        f.write(combined_md)
            except Exception as e:
                print(f'[Parser Error] Failed to parse presentation {filepath.name}: {e}')
                return []

        # 3. Word Documents (.docx)
        elif ext == '.docx':
            try:
                raw_chunks, combined_md = parse_docx_to_markdown(filepath)
                for ch in raw_chunks:
                    chunks.append({
                        'text': ch['text'],
                        'metadata': {
                            'source': str(filepath),
                            'filename': filepath.name,
                            'course': course_name,
                            'semester': sem_label,
                            'page': ch['page']
                        }
                    })
                if combined_md:
                    md_dest.parent.mkdir(parents=True, exist_ok=True)
                    with open(md_dest, 'w', encoding='utf-8') as f:
                        f.write(combined_md)
            except Exception as e:
                print(f'[Parser Error] Failed to parse DOCX {filepath.name}: {e}')
                return []

        # 4. ZIP Archives (.zip)
        elif ext == '.zip':
            try:
                extract_dir = filepath.parent / filepath.stem
                if not extract_dir.exists():
                    with zipfile.ZipFile(filepath, 'r') as zf:
                        for member in zf.infolist():
                            dest_file = (extract_dir / member.filename).resolve()
                            if not str(dest_file).startswith(str(extract_dir.resolve())):
                                continue
                            if member.is_dir():
                                dest_file.mkdir(parents=True, exist_ok=True)
                                continue
                            dest_file.parent.mkdir(parents=True, exist_ok=True)
                            with zf.open(member) as source, open(dest_file, 'wb') as target:
                                shutil.copyfileobj(source, target)
                
                zip_md_sections = [f'# Archive: {filepath.name}' + chr(10)]
                for sub_root, _, sub_files in os.walk(extract_dir):
                    for sub_f in sorted(sub_files):
                        if sub_f.startswith('.'):
                            continue
                        sub_path = Path(sub_root) / sub_f
                        sub_chunks = self.parse_file(sub_path, force=force)
                        for sc in sub_chunks:
                            chunks.append(sc)
                            zip_md_sections.append(f'## File: {sub_f}' + chr(10) + chr(10) + sc['text'])
                
                if zip_md_sections:
                    combined_md = (chr(10) + chr(10) + '---' + chr(10) + chr(10)).join(zip_md_sections)
                    md_dest.parent.mkdir(parents=True, exist_ok=True)
                    with open(md_dest, 'w', encoding='utf-8') as f:
                        f.write(combined_md)
            except Exception as e:
                print(f'[Parser Error] Failed to parse ZIP archive {filepath.name}: {e}')
                return []

        # 5. Plain text and Markdown
        elif ext in ('.txt', '.md'):
            try:
                with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                    text = f.read().strip()
                if text:
                    chunks.append({
                        'text': text,
                        'metadata': {
                            'source': str(filepath),
                            'filename': filepath.name,
                            'course': course_name,
                            'semester': sem_label,
                            'page': 1
                        }
                    })
                    md_dest.parent.mkdir(parents=True, exist_ok=True)
                    with open(md_dest, 'w', encoding='utf-8') as f:
                        f.write(text)
            except Exception as e:
                print(f'[Parser Error] Failed to parse text file {filepath.name}: {e}')
                return []

        # 6. HTML files
        elif ext in ('.html', '.htm'):
            try:
                with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                    soup = BeautifulSoup(f.read(), 'html.parser')
                for s in soup(['script', 'style', 'meta', 'noscript']):
                    s.extract()
                text = soup.get_text(separator=chr(10)).strip()
                if text:
                    chunks.append({
                        'text': text,
                        'metadata': {
                            'source': str(filepath),
                            'filename': filepath.name,
                            'course': course_name,
                            'semester': sem_label,
                            'page': 1
                        }
                    })
                    md_dest.parent.mkdir(parents=True, exist_ok=True)
                    with open(md_dest, 'w', encoding='utf-8') as f:
                        f.write(text)
            except Exception as e:
                print(f'[Parser Error] Failed to parse HTML {filepath.name}: {e}')
                return []

        # Update cache
        self.hashes[str(rel_path)] = curr_hash
        self._save_hashes()

        return chunks

    def parse_all(self, force=False):
        all_chunks = []
        supported_exts = {'.pdf', '.pptx', '.ppt', '.docx', '.zip', '.txt', '.md', '.html', '.htm'}
        
        for root_d, _, files in os.walk(self.output_dir):
            if '.dump' in root_d:
                continue
            for f in files:
                ext = Path(f).suffix.lower()
                if ext in supported_exts and not f.startswith('.'):
                    file_path = os.path.join(root_d, f)
                    chunks = self.parse_file(file_path, force=force)
                    all_chunks.extend(chunks)
        return all_chunks
